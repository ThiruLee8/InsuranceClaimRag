from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
import fitz

from app.core.logging import get_logger

logger = get_logger(__name__)


@dataclass
class ExtractedPage:
    page_number: int
    text: str


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage]
    page_count: int


@dataclass
class TextChunk:
    chunk_index: int
    page_number: int | None
    content: str


class DocumentProcessor:
    def extract_text(self, content: bytes, filename: str, content_type: str) -> ExtractedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf" or "pdf" in content_type:
            return self._extract_pdf(content)
        if suffix == ".docx" or "wordprocessingml" in content_type:
            return self._extract_docx(content)
        if suffix == ".txt" or content_type.startswith("text/"):
            return self._extract_txt(content)
        raise ValueError(f"Unsupported file type for extraction: {suffix or content_type}")

    def _extract_pdf(self, content: bytes) -> ExtractedDocument:
        pages: list[ExtractedPage] = []
        try:
            with fitz.open(stream=content, filetype="pdf") as doc:
                for idx, page in enumerate(doc, start=1):
                    text = page.get_text("text").strip()
                    if text:
                        pages.append(ExtractedPage(page_number=idx, text=text))
                page_count = doc.page_count
        except Exception:
            logger.warning("pymupdf_failed_falling_back_to_pypdf")
            reader = PdfReader(BytesIO(content))
            for idx, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append(ExtractedPage(page_number=idx, text=text))
            page_count = len(reader.pages)
        logger.info("pdf_extracted", page_count=page_count, pages_with_text=len(pages))
        return ExtractedDocument(pages=pages, page_count=page_count)

    def _extract_docx(self, content: bytes) -> ExtractedDocument:
        document = DocxDocument(BytesIO(content))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
        text = "\n".join(paragraphs)
        pages = [ExtractedPage(page_number=1, text=text)] if text else []
        logger.info("docx_extracted", character_count=len(text))
        return ExtractedDocument(pages=pages, page_count=1 if text else 0)

    def _extract_txt(self, content: bytes) -> ExtractedDocument:
        text = content.decode("utf-8", errors="ignore").strip()
        pages = [ExtractedPage(page_number=1, text=text)] if text else []
        logger.info("txt_extracted", character_count=len(text))
        return ExtractedDocument(pages=pages, page_count=1 if text else 0)

    def chunk_pages(
        self,
        extracted: ExtractedDocument,
        *,
        chunk_size: int,
        chunk_overlap: int,
    ) -> list[TextChunk]:
        if chunk_size <= 0:
            raise ValueError("chunk_size must be positive")
        if chunk_overlap < 0 or chunk_overlap >= chunk_size:
            raise ValueError("chunk_overlap must be >= 0 and < chunk_size")

        chunks: list[TextChunk] = []
        chunk_index = 0
        for page in extracted.pages:
            text = " ".join(page.text.split())
            if not text:
                continue
            start = 0
            while start < len(text):
                end = min(start + chunk_size, len(text))
                piece = text[start:end].strip()
                if piece:
                    chunks.append(
                        TextChunk(
                            chunk_index=chunk_index,
                            page_number=page.page_number,
                            content=piece,
                        )
                    )
                    chunk_index += 1
                if end >= len(text):
                    break
                start = max(0, end - chunk_overlap)
        logger.info("document_chunked", chunk_count=len(chunks))
        return chunks
